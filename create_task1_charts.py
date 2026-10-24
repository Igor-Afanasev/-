import math
from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Series, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

wb=Workbook()
info=wb.active;info.title='Условие'
info['A1']='Задание 1. Варианты 1 и 14'
info['A2']='Построить точечные диаграммы в декартовой системе координат'
info['A4']='Преобразование:';info['B4']='x = r·cos(t), y = r·sin(t)'
info['A6']='Функция 1';info['B6']='r(t) = 2 − 4sin(t)'
info['A7']='Интервал';info['B7']='0 ≤ t ≤ 2π; шаг 0,01'
info['A9']='Функция 2';info['B9']='r(t) = exp(sin(t)) − 2cos(4t) + sin((2t − π)/24)^5'
info['A10']='Интервал';info['B10']='−8π ≤ t ≤ 8π; шаг 0,05'
info['A12']='Примечание';info['B12']='Координаты рассчитаны численно; точечные диаграммы XY построены по столбцам x и y.'
info.column_dimensions['A'].width=23;info.column_dimensions['B'].width=90
info['A1'].font=Font(name='Calibri',size=17,bold=True,color='FFFFFF')
info['A1'].fill=PatternFill('solid',fgColor='163A60')
info.merge_cells('A1:B1')

def sheet(name,formula,left,right,step,kind):
    ws=wb.create_sheet(name)
    ws['A1']=name+' — точечная диаграмма'
    ws.merge_cells('A1:D1')
    ws['A2']=formula
    ws.merge_cells('A2:D2')
    for c,h in enumerate(['t (радианы)','r(t)','x = r cos(t)','y = r sin(t)'],1):
        cl=ws.cell(4,c,h);cl.font=Font(bold=True,color='FFFFFF');cl.fill=PatternFill('solid',fgColor='24496E');cl.alignment=Alignment(horizontal='center')
    steps=int(math.floor((right-left)/step+1e-8))
    vals=[left+i*step for i in range(steps+1)]
    if vals[-1] < right-1e-9: vals.append(right)
    for i,t in enumerate(vals,5):
        if kind==1:r=2-4*math.sin(t)
        else:r=math.exp(math.sin(t))-2*math.cos(4*t)+math.sin((2*t-math.pi)/24)**5
        for col,value in enumerate([t,r,r*math.cos(t),r*math.sin(t)],1):
            cl=ws.cell(i,col,value);cl.number_format='0.000000'
    ws.freeze_panes='A5'
    ws.auto_filter.ref=f'A4:D{4+len(vals)}'
    for x in 'ABCD':ws.column_dimensions[x].width=21
    ws.row_dimensions[1].height=27
    ws['A1'].font=Font(size=16,bold=True,color='163A60')
    ws['A2'].font=Font(size=11,italic=True)
    ws.sheet_view.showGridLines=False
    chart=ScatterChart()
    chart.title='Полярная функция в декартовых координатах'
    chart.style=13
    chart.x_axis.title='x'
    chart.y_axis.title='y'
    chart.scatterStyle='line'
    chart.legend=None
    chart.width=17;chart.height=14
    xref=Reference(ws,min_col=3,min_row=5,max_row=4+len(vals))
    yref=Reference(ws,min_col=4,min_row=5,max_row=4+len(vals))
    series=Series(yref,xref,title='r(t)')
    series.graphicalProperties.line.solidFill='0070C0'
    series.graphicalProperties.line.width=21000
    chart.series.append(series)
    ws.add_chart(chart,'F4')
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    return ws,len(vals)
sheet('График 1','r(t) = 2 − 4sin(t); t = 0…2π, шаг 0,01',0,2*math.pi,.01,1)
sheet('График 2','r(t) = e^sin(t) − 2cos(4t) + sin^5((2t−π)/24); t = −8π…8π, шаг 0,05',-8*math.pi,8*math.pi,.05,2)
for cell in ['A4','A6','A7','A9','A10','A12']:info[cell].font=Font(bold=True,color='24496E')
for row in info: 
 for cl in row:
  if cl.coordinate!='A1':cl.alignment=Alignment(vertical='top',wrap_text=True)
wb.active=1
out='Zadanie_1_varianty_1_i_14_grafiki.xlsx'
wb.save(out)
print(out)
