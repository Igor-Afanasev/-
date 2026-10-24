import math
from openpyxl import Workbook
from openpyxl.chart import SurfaceChart3D, Reference
from openpyxl.chart.axis import NumericAxis
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# z(x,y) = sqrt(5 |x|^3 cos^2(y)) - sqrt(3 |x| sin^4(y))
def f(x,y):
    return math.sqrt(5*abs(x)**3*math.cos(y)**2)-math.sqrt(3*abs(x)*math.sin(y)**4)

w=Workbook()
title=w.active;title.title='Условие'
title['A1']='Задание 2. График поверхности — варианты 1 и 16'
title['A3']='Функция z(x,y)'
title['B3']='SQRT(5*ABS(x)^3*COS(y)^2) - SQRT(3*ABS(x)*SIN(y)^4)'
title['A4']='Диапазоны'
title['B4']='x,y: от -5 до 5; шаг 0,1'
title['A5']='Построение'
title['B5']='Трёхмерная поверхность Excel по таблице значений z(x,y)'
title['A6']='Вертикальная ось z'
title['B6']='Минимум -5, максимум 25; значения выбраны так, чтобы охватить поверхность'
title['A8']='Замечание'
title['B8']='Таблица значений вычислена численно, формула исходного задания приведена выше.'
title.column_dimensions['A'].width=25; title.column_dimensions['B'].width=83
title['A1'].font=Font(size=16,bold=True,color='163A60')
for c in ('A3','A4','A5','A6','A8'):title[c].font=Font(bold=True,color='24496E')
surf=w.create_sheet('Поверхность')
surf['A1']='Задание 2 — 3D поверхность z(x,y)'
surf.merge_cells('A1:G1')
surf['A1'].font=Font(size=17,bold=True,color='17385C')
surf['A2']='z = √(5|x|³ cos²(y)) − √(3|x| sin⁴(y)); x,y ∈ [-5;5], шаг 0,1'
surf.merge_cells('A2:M2')
n=101
vals=[round(-5+i*.1,10) for i in range(n)]
surf['A4']='y \\ x'
for j,x in enumerate(vals,2):
    c=surf.cell(4,j,x);c.number_format='0.0'
for i,y in enumerate(vals,5):
    c=surf.cell(i,1,y);c.number_format='0.0'
    for j,x in enumerate(vals,2):
        c=surf.cell(i,j,f(x,y));c.number_format='0.0000'
for row in surf.iter_rows(min_row=4,max_row=4+n,min_col=1,max_col=n+1):
    for cell in row:
        if cell.row==4 or cell.column==1:
            cell.fill=PatternFill('solid',fgColor='1D4F78')
            cell.font=Font(color='FFFFFF',bold=True)
            cell.alignment=Alignment(horizontal='center')
surf.freeze_panes='B5'
surf.sheet_view.zoomScale=65
surf.column_dimensions['A'].width=11
for i in range(2,n+2):surf.column_dimensions[get_column_letter(i)].width=12
chart=SurfaceChart3D()
chart.title='Поверхность z(x,y)'
chart.style=26
chart.width=23;chart.height=16
chart.x_axis.title='x'
chart.y_axis.title='y'
chart.z_axis.title='z'
chart.z_axis.scaling.min=-5
chart.z_axis.scaling.max=25
chart.add_data(Reference(surf,min_col=2,max_col=n+1,min_row=4,max_row=n+4),titles_from_data=True,from_rows=True)
chart.set_categories(Reference(surf,min_col=1,min_row=5,max_row=n+4))
chart.legend=None
surf.add_chart(chart,'A108')
w.active=1
filename='Zadanie_2_varianty_1_i_16_poverhnost.xlsx'
w.save(filename)
print(filename)
